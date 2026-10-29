#!/usr/bin/env python3
"""Ajusta las tablas anchas del suplemento para que sean legibles en DOCX."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TARGETS = [
    ROOT / "entregables" / "en" / "material_suplementario.docx",
    ROOT / "entregables" / "es" / "material_suplementario_es.docx",
]
TOTAL_WIDTH = 14688


def _set_page(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)


def _set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for el in tc_pr.findall(qn("w:tcW")):
        tc_pr.remove(el)
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.insert_element_before(
        tc_w,
        "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
        "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
        "w:vAlign", "w:hideMark",
    )


def _widths(ncols: int) -> list[int]:
    if ncols >= 8:
        weights = [1.1, 0.9, 1.4] + [1.0] * (ncols - 3)
    elif ncols == 7:
        weights = [1.1, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0]
    else:
        weights = [1.0] * ncols
    base = TOTAL_WIDTH / sum(weights)
    out = [int(base * w) for w in weights]
    out[-1] += TOTAL_WIDTH - sum(out)
    return out


def _format_table(table) -> None:
    ncols = len(table.columns)
    if ncols == 0:
        return
    widths = _widths(ncols)
    size = 6.6 if ncols >= 18 else 6.8 if ncols >= 16 else 7.2 if ncols >= 14 else 7.6 if ncols >= 11 else 8.2

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for el in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(el)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")
    tbl_pr.insert_element_before(
        tbl_w, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
        "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
    )

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.insert_element_before(
            layout, "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
        )
    layout.set(qn("w:type"), "fixed")

    for el in tbl_pr.findall(qn("w:tblCellMar")):
        tbl_pr.remove(el)
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "36" if side in ("left", "right") else "0")
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tbl_pr.insert_element_before(
        margins, "w:tblLook", "w:tblCaption", "w:tblDescription",
    )

    grid = tbl.tblGrid
    for col in list(grid.findall(qn("w:gridCol"))):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ridx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.insert_element_before(
                OxmlElement("w:cantSplit"), "w:trHeight", "w:tblHeader",
                "w:tblCellSpacing", "w:jc", "w:hidden",
            )
        if ridx == 0 and tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.insert_element_before(
                OxmlElement("w:tblHeader"), "w:tblCellSpacing", "w:jc", "w:hidden",
            )
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                _set_cell_width(cell, widths[idx])

    for ridx, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(size)
                    if ridx == 0:
                        run.font.bold = True


def _format(path: Path) -> None:
    doc = Document(path)
    for section in doc.sections:
        _set_page(section)
    for table in doc.tables:
        _format_table(table)
    doc.save(path)
    print(f"{path.resolve().relative_to(ROOT)}: {len(doc.tables)} tablas")


def main(argv: list[str]) -> int:
    targets = [Path(arg) for arg in argv[1:]] or DEFAULT_TARGETS
    for target in targets:
        if not target.exists():
            raise FileNotFoundError(target)
        _format(target)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
