"""Generate the final figure PDF and table workbook from the manuscript sources."""

from __future__ import annotations

import re
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
MANUSCRIPT = ROOT / "manuscrito" / "en" / "manuscrito.qmd"
SUPPLEMENT = ROOT / "manuscrito" / "en" / "material_suplementario.qmd"
MAIN_DOCX = ROOT / "entregables" / "manuscrito.docx"
SUPP_DOCX = ROOT / "entregables" / "material_suplementario.docx"
FIGURES_PDF = ROOT / "entregables" / "figuras_maestro.pdf"
TABLES_XLSX = ROOT / "entregables" / "tablas_maestro.xlsx"
ARIAL = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
ARIAL_BOLD = Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")


def _plain(text: str) -> str:
    text = re.sub(r"!\[(.*?)\]\(.*?\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    for old, new in (
        (r"\alpha", "α"),
        (r"\beta", "β"),
        (r"\Delta", "Δ"),
        (r"\approx", "≈"),
        (r"\geq", "≥"),
        (r"\leq", "≤"),
        (r"\times", "×"),
        (r"\pm", "±"),
        (r"\to", "→"),
        (r"\cdot", "·"),
        (r"\%", "%"),
        (r"\,", " "),
        (r"\;", " "),
    ):
        text = text.replace(old, new)
    text = text.replace("$", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _figures_from_qmd(path: Path, main: bool) -> list[tuple[str, Path]]:
    figures: list[tuple[str, Path]] = []
    counter = 0
    for line in path.read_text(encoding="utf-8").splitlines():
        match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if not match:
            continue
        caption = _plain(match.group(1))
        image = (path.parent / match.group(2)).resolve()
        if main:
            counter += 1
            caption = f"Figure {counter}. {caption}"
        figures.append((caption, image))
    return figures


def build_figures() -> None:
    figures = _figures_from_qmd(MANUSCRIPT, main=True) + _figures_from_qmd(SUPPLEMENT, main=False)
    doc = fitz.open()
    toc: list[list[object]] = []
    for caption, image in figures:
        if not image.exists():
            raise FileNotFoundError(image)
        page = doc.new_page(width=595, height=842)
        page.insert_font(fontname="F0", fontfile=str(ARIAL))
        page.insert_font(fontname="F1", fontfile=str(ARIAL_BOLD))
        page.insert_text((40, 34), caption.split(".")[0] + ".", fontsize=9, fontname="F1")
        with Image.open(image) as img:
            iw, ih = img.size
        area = fitz.Rect(40, 48, 555, 700)
        scale = min(area.width / iw, area.height / ih)
        iw, ih = iw * scale, ih * scale
        image_rect = fitz.Rect(
            area.x0 + (area.width - iw) / 2,
            area.y0 + (area.height - ih) / 2,
            area.x0 + (area.width + iw) / 2,
            area.y0 + (area.height + ih) / 2,
        )
        page.insert_image(image_rect, filename=str(image), keep_proportion=True)
        caption_rect = fitz.Rect(40, 704, 555, 812)
        for font_size in (8.2, 7.8, 7.4, 7.0, 6.6, 6.2):
            if page.insert_textbox(
                caption_rect,
                caption,
                fontsize=font_size,
                fontname="F0",
                lineheight=1.12,
                align=0,
            ) >= 0:
                break
        page.insert_text(
            (40, 826),
            "Master figures | Incendios_PY v2.1.0",
            fontsize=7,
            fontname="F0",
            color=(0.35, 0.35, 0.35),
        )
        label = caption.split(".")[0]
        toc.append([1, f"{label} — {caption.split('. ', 1)[-1][:80]}", page.number + 1])
    doc.set_toc(toc)
    doc.set_metadata({"title": "Master figures", "author": "Martín Sánchez Insfrán"})
    doc.save(FIGURES_PDF, garbage=4, deflate=True)
    doc.close()
    print(f"figuras_maestro.pdf: {len(figures)} figures")


def _row_values(table: Table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([re.sub(r"\s+", " ", cell.text.replace("\xa0", " ")).strip() for cell in row.cells])
    return rows


def _paragraph_text_math(paragraph: Paragraph) -> str:
    parts: list[str] = []
    for node in paragraph._p.iter():
        if node.tag in {qn("w:t"), qn("m:t")} and node.text:
            parts.append(node.text)
    return "".join(parts)


def _caption_from_paragraph(paragraph: Paragraph) -> str:
    return re.sub(r"\s+", " ", _paragraph_text_math(paragraph).replace("\xa0", " ")).strip()


def _table_caption(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^Table\s+(S?\d+[a-z]?)\s*[.:]\s*", "", text)
    return text


def _base_name(caption: str) -> str:
    match = re.match(r"^Table\s+(S?\d+[a-z]?)", caption)
    if not match:
        return "T"
    return "T" + match.group(1) if not match.group(1).startswith("S") else "T" + match.group(1)


def _main_tables() -> list[tuple[str, list[list[str]]]]:
    doc = Document(MAIN_DOCX)
    tables: list[tuple[str, list[list[str]]]] = []
    for table in doc.tables:
        if not table.cell(0, 0).paragraphs:
            continue
        caption = _caption_from_paragraph(table.cell(0, 0).paragraphs[0])
        if not caption.startswith("Table "):
            continue
        nested = table.cell(0, 0).tables
        if not nested:
            continue
        tables.append((caption, _row_values(nested[0])))
    return tables


def _supplement_tables() -> list[tuple[str, list[list[str]]]]:
    doc = Document(SUPP_DOCX)
    tables: list[tuple[str, list[list[str]]]] = []
    caption = ""
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = _caption_from_paragraph(Paragraph(child, doc))
            if text.startswith("Table S"):
                caption = text
        elif child.tag.endswith("}tbl"):
            if not caption:
                continue
            tables.append((caption, _row_values(Table(child, doc))))
    return tables


def _unique_name(caption: str, used: set[str]) -> str:
    base = _base_name(caption)
    if base not in used:
        used.add(base)
        return base
    suffix = ord("b")
    while f"{base}{chr(suffix)}" in used:
        suffix += 1
    name = f"{base}{chr(suffix)}"
    used.add(name)
    return name


def _style_sheet(ws, caption: str, rows: list[list[str]], source: str) -> None:
    width = max(len(row) for row in rows)
    last_col = get_column_letter(width)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    title = ws.cell(row=1, column=1, value=caption)
    title.font = Font(bold=True, size=11)
    title.alignment = Alignment(wrap_text=True, vertical="top")
    title.fill = PatternFill("solid", fgColor="D9E2F3")
    ws.row_dimensions[1].height = 42
    for col, value in enumerate(rows[0], 1):
        cell = ws.cell(row=3, column=col, value=value)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for row_idx, row in enumerate(rows[1:], 4):
        for col, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{last_col}{max(3, len(rows) + 2)}"
    ws.sheet_view.showGridLines = False
    for col in range(1, width + 1):
        values = [str(rows[r][col - 1]) if col <= len(rows[r]) else "" for r in range(min(len(rows), 80))]
        ws.column_dimensions[get_column_letter(col)].width = min(48, max(10, max(len(v) for v in values) + 2))
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = source


def build_tables() -> None:
    wb = Workbook()
    index = wb.active
    index.title = "Indice"
    index.sheet_view.showGridLines = False
    index.merge_cells("A1:G1")
    index["A1"] = "Master tables | Incendios_PY v2.1.0"
    index["A1"].font = Font(bold=True, size=14)
    headers = ["Hoja", "Tabla", "Título", "Fuente", "Filas", "Columnas"]
    for col, value in enumerate(headers, 1):
        cell = index.cell(row=3, column=col, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    index.freeze_panes = "A4"
    index.auto_filter.ref = f"A3:F{3 + len(_main_tables()) + len(_supplement_tables())}"
    used: set[str] = set()
    row_idx = 4
    for source, tables in (("manuscrito", _main_tables()), ("material_suplementario", _supplement_tables())):
        for caption, rows in tables:
            name = _unique_name(caption, used)
            sheet = wb.create_sheet(name)
            _style_sheet(sheet, caption, rows, source)
            values = [name, caption.split(".")[0], _table_caption(caption), source, len(rows) - 1, max(len(r) for r in rows)]
            for col, value in enumerate(values, 1):
                index.cell(row=row_idx, column=col, value=value)
            link = index.cell(row=row_idx, column=1)
            link.hyperlink = f"#'{name}'!A1"
            link.style = "Hyperlink"
            row_idx += 1
    widths = [12, 12, 70, 24, 8, 9]
    for col, width in enumerate(widths, 1):
        index.column_dimensions[get_column_letter(col)].width = width
    for row in index.iter_rows(min_row=4, max_row=row_idx - 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(TABLES_XLSX)
    print(f"tablas_maestro.xlsx: {row_idx - 4} tables")


if __name__ == "__main__":
    build_figures()
    build_tables()
