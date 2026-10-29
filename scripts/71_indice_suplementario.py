"""Create a navigable supplement copy with a clickable contents index."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "entregables" / "material_suplementario.docx"
TARGET = ROOT / "entregables" / "material_suplementario_navegable.docx"


def _bookmark_id(document: Document) -> int:
    ids = [int(node.get(qn("w:id"))) for node in document.element.body.iter(qn("w:bookmarkStart"))]
    return max(ids, default=0) + 1


def _bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is not None:
        p_pr.addnext(start)
    else:
        paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _hyperlink(paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.extend([color, size, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")[:32]


def _label(text: str, limit: int = 118) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    return text if len(text) <= limit else text[: limit - 3].rstrip() + "..."


def _new_after(reference, document: Document):
    paragraph = document.add_paragraph()
    reference._p.addnext(paragraph._p)
    return paragraph


def _style_entry(paragraph, text: str, anchor: str, bold: bool = False, indent: float = 0.22) -> None:
    paragraph.paragraph_format.left_indent = Pt(indent * 72)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    if bold:
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
        return
    bullet = paragraph.add_run("• ")
    bullet.font.size = Pt(9)
    _hyperlink(paragraph, text, anchor)


def _find_contents(document: Document):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "Supplementary contents":
            return paragraph
    raise ValueError("No se encontró 'Supplementary contents'")


def _remove_old_bullets(document: Document, contents) -> None:
    index = next(i for i, p in enumerate(document.paragraphs) if p._p is contents._p)
    for paragraph in document.paragraphs[index + 1 :]:
        text = paragraph.text.strip()
        if text.startswith("Supplementary Note S"):
            break
        if text.startswith("•"):
            paragraph._p.getparent().remove(paragraph._p)


def _entries(document: Document):
    notes = []
    tables = []
    figures = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if re.match(r"^Supplementary Note S\d+\.", text):
            notes.append((text, paragraph))
        elif re.match(r"^Table S\d+[a-z]?(?: \(continued\))?\.", text):
            tables.append((text, paragraph))
        elif re.match(r"^Figure S\d+\.", text):
            figures.append((text, paragraph))
    return notes, tables, figures


def _add_bookmarks(document: Document, groups):
    bookmark_id = _bookmark_id(document)
    used: set[str] = set()
    occurrences: dict[str, int] = {}
    indexed = []
    for group, kind in zip(groups, ("note", "table", "figure")):
        for text, paragraph in group:
            if kind == "note":
                match = re.match(r"Supplementary Note (S\d+)", text)
                base = match.group(1) if match else "note"
                label = re.sub(r"^Supplementary Note ", "Note ", text)
            elif kind == "table":
                match = re.match(r"Table (S\d+[a-z]?)", text)
                base = match.group(1) if match else "table"
                label = text
            else:
                match = re.match(r"Figure (S\d+)", text)
                base = match.group(1) if match else "figure"
                label = text
            occurrence_key = (kind, base)
            occurrences[occurrence_key] = occurrences.get(occurrence_key, 0) + 1
            suffix = "" if occurrences[occurrence_key] == 1 else f"_{occurrences[occurrence_key]}"
            name = f"idx_{kind}_{_slug(base)}{suffix}"[:40]
            while name in used:
                name += "x"
            used.add(name)
            _bookmark(paragraph, name, bookmark_id)
            bookmark_id += 1
            if occurrences[occurrence_key] > 1:
                label = f"{label} [{occurrences[occurrence_key]}]"
            indexed.append((kind, label, name))
    return indexed


def build() -> None:
    document = Document(SOURCE)
    contents = _find_contents(document)
    _remove_old_bullets(document, contents)
    groups = _entries(document)
    indexed = _add_bookmarks(document, groups)

    intro = _new_after(contents, document)
    intro.paragraph_format.space_before = Pt(4)
    intro.paragraph_format.space_after = Pt(4)
    intro.paragraph_format.line_spacing = 1.0
    run = intro.add_run("Click any entry to jump to the corresponding note, table or figure.")
    run.italic = True
    run.font.size = Pt(8.5)

    labels = {"note": "Supplementary notes", "table": "Tables", "figure": "Figures"}
    cursor = intro
    current = None
    for kind in ("note", "table", "figure"):
        entries = [item for item in indexed if item[0] == kind]
        heading = _new_after(cursor, document)
        heading.paragraph_format.space_before = Pt(7)
        heading.paragraph_format.space_after = Pt(2)
        run = heading.add_run(f"{labels[kind]} ({len(entries)} entries)")
        run.bold = True
        run.font.size = Pt(10)
        cursor = heading
        for _, label, anchor in entries:
            item = _new_after(cursor, document)
            _style_entry(item, _label(label), anchor)
            cursor = item

    document.save(TARGET)
    print(f"Guardado: {TARGET.relative_to(ROOT)}")
    print(f"Notas: {len(groups[0])}; tablas: {len(groups[1])}; figuras: {len(groups[2])}")


if __name__ == "__main__":
    build()
