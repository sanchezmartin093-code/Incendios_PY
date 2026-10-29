"""Check cross-references, assets and submission metadata before release."""

from __future__ import annotations

import re
from pathlib import Path

import fitz
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
MAIN = ROOT / "manuscrito" / "en" / "manuscrito.qmd"
SUPP = ROOT / "manuscrito" / "en" / "material_suplementario.qmd"
HIGHLIGHTS = ROOT / "highlights.md"
OUTPUTS = [
    ROOT / "entregables" / "manuscrito.docx",
    ROOT / "entregables" / "material_suplementario.docx",
    ROOT / "entregables" / "highlights.docx",
    ROOT / "entregables" / "cover_letter.docx",
    ROOT / "entregables" / "figuras_maestro.pdf",
    ROOT / "entregables" / "tablas_maestro.xlsx",
]


def _check_qmd(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    definitions = set(re.findall(r"\{#((?:tbl|fig)-[\w-]+)", text))
    references = set(re.findall(r"@((?:tbl|fig)-[\w-]+)", text))
    missing = sorted(references - definitions)
    errors = [f"{path.name}: referencia sin destino @{item}" for item in missing]
    for image in re.findall(r"!\[.*?\]\((.*?)\)", text):
        if not (path.parent / image).resolve().exists():
            errors.append(f"{path.name}: imagen ausente {image}")
    return errors


def main() -> int:
    errors: list[str] = []
    for path in (MAIN, SUPP):
        errors.extend(_check_qmd(path))

    main_text = MAIN.read_text(encoding="utf-8")
    if "Incendios_PY_data_v2.0.1.zip" not in main_text:
        errors.append("manuscrito: falta el asset de datos v2.0.1")
    if "Incendios_PY_data_v2.1.0.zip" in main_text:
        errors.append("manuscrito: referencia a un asset de datos v2.1.0 inexistente")
    if "#fig-graphical" in main_text:
        errors.append("manuscrito: graphical abstract numerado dentro del texto")
    abstract = main_text.split("**ABSTRACT**", 1)[1].split("# 1. Introduction", 1)[0]
    words = len(re.sub(r"[*_`$]", "", abstract).split())
    if not 200 <= words <= 300:
        errors.append(f"abstract: {words} palabras")
    if "reliability diagram" in main_text:
        errors.append("manuscrito: referencia a reliability diagram inexistente")

    supp_text = SUPP.read_text(encoding="utf-8")
    if "Figures S1–S10" not in supp_text:
        errors.append("suplemento: el índice no incluye Figures S1–S10")
    if re.search(r"Table (?:10|4)\b", supp_text):
        errors.append("suplemento: referencia huérfana a Table 10/Table 4")
    if "vapour" in supp_text.lower():
        errors.append("suplemento: ortografía británica vapour")

    highlights = []
    for line in HIGHLIGHTS.read_text(encoding="utf-8").splitlines():
        line = line.strip().lstrip("-• ").strip()
        if line and not line.startswith("#"):
            highlights.append(line)
    if not 3 <= len(highlights) <= 5:
        errors.append(f"highlights: {len(highlights)} viñetas")
    for item in highlights:
        if len(item) > 85:
            errors.append(f"highlights: {len(item)} caracteres: {item[:40]}")

    for output in OUTPUTS:
        if not output.exists():
            errors.append(f"salida ausente: {output.relative_to(ROOT)}")

    figures = fitz.open(ROOT / "entregables" / "figuras_maestro.pdf")
    if figures.page_count != 15 or len(figures.get_toc()) != 15:
        errors.append(f"figuras_maestro: {figures.page_count} páginas, {len(figures.get_toc())} marcadores")
    figures.close()

    workbook = openpyxl.load_workbook(ROOT / "entregables" / "tablas_maestro.xlsx", read_only=True)
    if "Indice" not in workbook.sheetnames or len(workbook.sheetnames) != 50:
        errors.append(f"tablas_maestro: {len(workbook.sheetnames)} hojas")
    workbook.close()

    print(f"abstract: {words} palabras")
    print(f"highlights: {len(highlights)} viñetas")
    print(f"figuras_maestro: {fitz.open(ROOT / 'entregables' / 'figuras_maestro.pdf').page_count} páginas")
    print(f"tablas_maestro: {len(openpyxl.load_workbook(ROOT / 'entregables' / 'tablas_maestro.xlsx', read_only=True).sheetnames)} hojas")
    if errors:
        for error in errors:
            print("ERROR:", error)
        return 1
    print("check: OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
